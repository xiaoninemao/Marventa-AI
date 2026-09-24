from __future__ import annotations

import re
from docx import Document as DocxDocument
from app.engines.market_insight.models import ParsedDocument, Section, CodeBlock


def parse_docx(file_path: str) -> ParsedDocument:
    doc = DocxDocument(file_path)
    paragraphs = [p.text for p in doc.paragraphs]
    content = "\n".join(paragraphs)
    title = _extract_docx_title(doc, file_path)
    sections = _extract_sections_from_paragraphs(paragraphs)
    code_blocks = _extract_code_blocks(content)
    tech_stack = _detect_tech_stack(content)
    features = _extract_features_from_sections(sections)

    return ParsedDocument(
        title=title,
        source_type="docx",
        sections=sections,
        code_blocks=code_blocks,
        tech_stack=tech_stack,
        features=features,
        raw_text=content,
    )


def _extract_docx_title(doc: DocxDocument, file_path: str) -> str:
    import os
    for para in doc.paragraphs:
        text = para.text.strip()
        if text and para.style and "Heading" in para.style.name and len(text) < 150:
            return text
    for para in doc.paragraphs:
        text = para.text.strip()
        if text and len(text) < 150:
            return text
    return os.path.splitext(os.path.basename(file_path))[0]


def _extract_sections_from_paragraphs(paragraphs: list[str]) -> list[Section]:
    sections: list[Section] = []
    current_section: Section | None = None

    for text in paragraphs:
        stripped = text.strip()
        if not stripped:
            continue

        looks_like_heading = (
            len(stripped) < 120
            and not stripped.endswith(".")
            and (stripped[0].isupper() or stripped[0].isdigit())
        )

        if looks_like_heading and current_section and current_section.content:
            sections.append(current_section)
            current_section = Section(heading=stripped, level=1, content="")
        elif looks_like_heading and not current_section:
            current_section = Section(heading=stripped, level=1, content="")
        elif current_section:
            current_section.content += stripped + "\n"
        else:
            current_section = Section(heading="Preamble", level=0, content=stripped)

    if current_section:
        sections.append(current_section)

    return sections


def _extract_code_blocks(content: str) -> list[CodeBlock]:
    pattern = re.compile(r"```(\w*)\n(.*?)```", re.DOTALL)
    return [
        CodeBlock(language=m.group(1) or "", code=m.group(2).strip())
        for m in pattern.finditer(content)
    ]


TECH_KEYWORDS = {
    "python", "javascript", "typescript", "go", "golang", "rust", "java",
    "kotlin", "swift", "react", "vue", "angular", "next.js", "nextjs",
    "nuxt", "node.js", "nodejs", "django", "flask", "fastapi", "spring",
    "postgresql", "mysql", "mongodb", "redis", "elasticsearch",
    "docker", "kubernetes", "k8s", "aws", "gcp", "azure", "terraform",
    "graphql", "rest", "grpc", "kafka", "rabbitmq", "nginx",
    "pytorch", "tensorflow", "langchain", "huggingface", "llm",
    "openai", "anthropic", "claude", "gpt", "embedding",
}


def _detect_tech_stack(content: str) -> list[str]:
    lower = content.lower()
    found = set()
    for kw in TECH_KEYWORDS:
        if kw in lower:
            found.add(kw)
    return sorted(found)


def _extract_features_from_sections(sections: list[Section]) -> list[str]:
    features: list[str] = []
    for sec in sections:
        for line in sec.content.split("\n"):
            stripped = line.strip()
            if stripped.startswith(("- ", "* ", "• ")):
                features.append(stripped.lstrip("- *•").strip())
    return features[:20]
